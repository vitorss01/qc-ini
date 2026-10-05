# -*- coding: utf-8 -*-
"""Gerador de DB_SEAC SINTETICO para os testes do ETL de recebimento (qa_etl_recebimento.py).

Cria um .xlsx com UMA tabela Excel cujo nome e o CFG TABELA_ORIGEM do produto
(tbHematologia / tbBioquimica) e com as colunas EXATAMENTE como o recebimento espera:
as colunas da tblDB_Recebimento da copia em teste, menos ID_REGISTRO e RECEBIDO_EM
(o cabecalho e lido EM TEMPO DE EXECUCAO pelo teste e passado aqui -- nunca fixo no codigo).

Tipos gravados = os de pq/SEAC_ORIGEM.m: ID_ORIGEM, ITEM_ID e NIVEL inteiros; DATA data;
DATA_HORA data-hora; VALOR numero; o resto texto. Os valores CLONADOS do modelo (linha real
lida por COM: inteiros chegam como float, datas com fuso) sao convertidos para esses tipos.
Os valores passados explicitamente em cada linha sao gravados COMO VIERAM -- e assim que o
teste grava um ID_ORIGEM em texto (' 9000010 '), fracionario (9000009.4) ou nulo.

Uso pelo teste:
    escritas = gerar(caminho_xlsx, produto, cabecalho, modelo, linhas, tabela='tbBioquimica')
      modelo ... dict coluna -> valor (linha REAL do recebimento usada como base)
      linhas ... lista de dicts com os campos que mudam; a chave '_modelo' troca a linha-base
                 daquela linha (ex.: o molde do nivel 2); chaves iniciadas por '_' nao sao gravadas
      devolve .. a lista do que foi gravado (dict coluna -> valor), na ordem das linhas

Autoteste sem Excel:  python gerar_seac_sintetico.py --autoteste <saida.xlsx>
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TABELAS = {'HEMATOLOGIA': 'tbHematologia', 'BIOQUIMICA': 'tbBioquimica'}
ABAS = {'tbHematologia': 'DADOS_HEMATOLOGIA', 'tbBioquimica': 'DADOS_BIOQUIMICA'}

# pq/SEAC_ORIGEM.m (TiposConhecidos)
INTEIRAS = {'ID_ORIGEM', 'ITEM_ID', 'NIVEL'}
DATAS = {'DATA'}
DATAS_HORA = {'DATA_HORA'}
NUMEROS = {'VALOR'}
TEXTOS = {'EQUIPAMENTO', 'MATRIZ', 'BLOCO', 'LOTE', 'NIVEL_DESC', 'ANALITO', 'PRIMEIRO_DO_DIA', 'FLAG',
          'ID_AMOSTRA', 'UNIDADE'}
PROIBIDAS = {'ID_REGISTRO', 'RECEBIDO_EM'}       # criadas pelo recebimento, nunca vem da origem
OBRIGATORIAS = {'ID_ORIGEM', 'DATA', 'DATA_HORA', 'VALOR', 'NIVEL', 'ANALITO'}


def _sem_fuso(v):
    if isinstance(v, dt.datetime) and v.tzinfo is not None:
        return v.replace(tzinfo=None)
    return v


def _dt_puro(v):
    """pywintypes.datetime -> datetime.datetime ingenuo (openpyxl nao conhece o tipo do COM)."""
    v = _sem_fuso(v)
    if isinstance(v, dt.datetime):
        return dt.datetime(v.year, v.month, v.day, v.hour, v.minute, v.second, v.microsecond)
    return v


def tipar(coluna, v):
    """Valor CLONADO do modelo -> tipo da origem (pq/SEAC_ORIGEM.m)."""
    v = _dt_puro(v)
    if v is None or (isinstance(v, str) and v == ''):
        return None
    if coluna in INTEIRAS:
        if isinstance(v, float) and v.is_integer():
            return int(v)
        if isinstance(v, str) and v.strip().isdigit():
            return int(v.strip())
        return v
    if coluna in DATAS:
        if isinstance(v, dt.datetime):
            return v.date()
        return v
    if coluna in DATAS_HORA:
        if isinstance(v, dt.date) and not isinstance(v, dt.datetime):
            return dt.datetime(v.year, v.month, v.day)
        return v
    if coluna in NUMEROS:
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return float(v)
        return v
    if coluna in TEXTOS:
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return str(v)
        if isinstance(v, dt.datetime):
            return v.strftime('%d/%m/%Y %H:%M:%S')
        return v
    return v


def tabela_do_produto(produto):
    p = str(produto).upper().replace('Í', 'I')
    for k, v in TABELAS.items():
        if p.startswith(k[:3]):
            return v
    raise ValueError(f'produto desconhecido: {produto!r}')


def montar_linhas(cabecalho, modelo, linhas):
    """O que sera gravado: modelo (tipado) + campos da linha (como vieram)."""
    escritas = []
    for ln in linhas:
        base = ln.get('_modelo') or modelo or {}
        out = {}
        for c in cabecalho:
            if c in ln:
                out[c] = _dt_puro(ln[c])
            else:
                out[c] = tipar(c, base.get(c))
        escritas.append(out)
    return escritas


def gerar(caminho_xlsx, produto, cabecalho, modelo, linhas, tabela=None):
    """Cria o DB_SEAC sintetico. Devolve a lista do que foi gravado (uma entrada por linha)."""
    tabela = (tabela or tabela_do_produto(produto)).strip()
    cab = [str(c) for c in cabecalho]
    if len(set(cab)) != len(cab):
        raise ValueError(f'cabecalho com coluna repetida: {cab}')
    if PROIBIDAS & set(cab):
        raise ValueError(f'cabecalho nao pode ter {sorted(PROIBIDAS & set(cab))} (o recebimento cria)')
    if not OBRIGATORIAS <= set(cab):
        raise ValueError(f'cabecalho sem as colunas obrigatorias {sorted(OBRIGATORIAS - set(cab))}')
    if not linhas:
        raise ValueError('pelo menos uma linha sintetica (tabela Excel nao pode ficar sem corpo)')
    escritas = montar_linhas(cab, modelo, linhas)

    wb = Workbook()
    ws = wb.active
    ws.title = ABAS.get(tabela, 'DADOS')
    for j, c in enumerate(cab, start=1):
        ws.cell(row=1, column=j, value=c)
    for i, out in enumerate(escritas, start=2):
        for j, c in enumerate(cab, start=1):
            v = out[c]
            if v is None:
                continue                                  # nulo explicito = celula vazia
            cel = ws.cell(row=i, column=j, value=v)
            if isinstance(v, dt.datetime):
                cel.number_format = 'dd/mm/yyyy hh:mm:ss'
            elif isinstance(v, dt.date):
                cel.number_format = 'dd/mm/yyyy'
    ref = f'A1:{get_column_letter(len(cab))}{len(escritas) + 1}'
    t = Table(displayName=tabela, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)
    pasta = os.path.dirname(os.path.abspath(caminho_xlsx))
    os.makedirs(pasta, exist_ok=True)
    tmp = os.path.abspath(caminho_xlsx) + '.tmp.xlsx'
    wb.save(tmp)
    os.replace(tmp, caminho_xlsx)                         # o Power Query nunca le arquivo pela metade
    return escritas


def ler(caminho_xlsx, tabela):
    """Le de volta a tabela gravada (conferencia do proprio gerador; sem Excel)."""
    wb = load_workbook(caminho_xlsx)
    for ws in wb.worksheets:
        if tabela in ws.tables:
            rng = ws[ws.tables[tabela].ref]
            cab = [c.value for c in rng[0]]
            return [dict(zip(cab, [c.value for c in lin])) for lin in rng[1:]]
    raise KeyError(f'tabela {tabela} nao encontrada em {caminho_xlsx}')


def _autoteste(saida):
    cab = ['EQUIPAMENTO', 'LOTE', 'NIVEL', 'NIVEL_DESC', 'DATA', 'DATA_HORA', 'ANALITO', 'VALOR', 'FLAG',
           'ID_ORIGEM', 'ID_AMOSTRA', 'ITEM_ID', 'UNIDADE']
    modelo = {'EQUIPAMENTO': 'XN1000', 'LOTE': 8974.0, 'NIVEL': 1.0, 'NIVEL_DESC': 'NIVEL 1',
              'DATA': dt.datetime(2026, 5, 3, tzinfo=dt.timezone.utc), 'DATA_HORA': dt.datetime(2026, 5, 3, 7, 1),
              'ANALITO': 'WBC', 'VALOR': 7.12, 'FLAG': None, 'ID_ORIGEM': 123.0, 'ID_AMOSTRA': 'A1',
              'ITEM_ID': 55.0, 'UNIDADE': '10^3/uL', 'RECEBIDO_EM': dt.datetime(2026, 5, 3)}
    dia = dt.date(2026, 5, 10)
    linhas = [{'ID_ORIGEM': 9000001, 'ITEM_ID': 9000001, 'DATA': dia, 'DATA_HORA': dt.datetime(2026, 5, 10, 8, 0)},
              {'ID_ORIGEM': ' 9000010 ', 'ITEM_ID': 9000010, 'DATA': dia, 'DATA_HORA': dt.datetime(2026, 5, 10, 8, 1)},
              {'ID_ORIGEM': 9000009.4, 'ITEM_ID': 9000109, 'DATA': dia, 'DATA_HORA': None},
              {'ID_ORIGEM': None, 'ITEM_ID': 9000006, 'DATA': dia, 'DATA_HORA': dt.datetime(2026, 5, 10, 9, 0)}]
    esc = gerar(saida, 'Hematologia', cab, modelo, linhas)
    lido = ler(saida, 'tbHematologia')
    erros = []
    if len(lido) != len(esc):
        erros.append(('linhas', len(esc), len(lido)))
    for e, l in zip(esc, lido):
        for c in cab:
            a, b = e[c], l[c]
            if isinstance(a, dt.date) and not isinstance(a, dt.datetime):
                a = dt.datetime(a.year, a.month, a.day)
            if a != b:
                erros.append((c, a, b))
    tipos = {c: type(esc[0][c]).__name__ for c in cab}
    print({'arquivo': saida, 'linhas': len(lido), 'tipos_linha1': tipos, 'divergencias': erros})
    return not erros


if __name__ == '__main__':
    if len(sys.argv) >= 3 and sys.argv[1] == '--autoteste':
        sys.exit(0 if _autoteste(os.path.abspath(sys.argv[2])) else 1)
    print(__doc__)
