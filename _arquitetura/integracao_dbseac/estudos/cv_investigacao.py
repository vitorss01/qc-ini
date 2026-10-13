# -*- coding: utf-8 -*-
"""cv_investigacao.py -- investigacao do CV alto do CIQ (somente leitura, sobre uma COPIA do .xlsm).

Uso:  python cv_investigacao.py <Bioquimica|Hematologia> <copia.xlsm> <saida_prefixo> [--meses 6]

Mesma populacao do Sigma (ADR-071): tblCQ_Final com PARTICIPA_ESTATISTICA = SIM, janela de N meses ancorada no
ultimo resultado, por EQUIPAMENTO x ANALITO x NIVEL, lotes separados; CV pooled = raiz(soma((n-1) CV^2)/soma(n-1))
so com lotes n >= 20 (mIncerteza.IncertezaCIQ). Para cada par aponta os resultados SUSPEITOS, que inflam o CV:
  - valor <= 0 (controle nao tem resultado zero/negativo);
  - fora de 5 DP robustos do proprio lote (|x - mediana| / (1,4826 x MAD) > 5; com MAD = 0, desvio medio absoluto).
e recalcula o CV pooled SEM eles, para o RT ver o efeito. NADA e inativado: a decisao e do laboratorio, pela aba
Inativar (ID + analito + motivo), com a justificativa tecnica.
"""
import csv
import math
import statistics as st
import sys
from collections import defaultdict
from datetime import datetime, timedelta

import openpyxl

N_LOTE_MIN = 20
Z_ROB = 5.0


def ler(caminho):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = next(w for w in wb.worksheets if w.title.startswith('Principal'))
    linhas = ws.iter_rows(min_row=6, values_only=True)
    cab = [str(c) for c in next(linhas)]
    ix = {c: i for i, c in enumerate(cab)}
    dados = []
    for r in linhas:
        if r[0] in (None, ''):
            continue
        if str(r[ix['PARTICIPA_ESTATISTICA']]).upper() != 'SIM':
            continue
        v, d = r[ix['RESULTADO']], r[ix['DATA']]
        if not isinstance(v, (int, float)) or not isinstance(d, datetime):
            continue
        dados.append({'id': str(r[0]), 'eq': str(r[ix['EQUIPAMENTO']] or ''), 'an': str(r[ix['ANALITO']] or ''),
                      'nv': r[ix['NIVEL']], 'lote': str(r[ix['LOTE']] or ''), 'data': d,
                      'dh': r[ix['DATA_HORA']], 'v': float(v)})
    est = next(w for w in wb.worksheets if w.title == 'Estatística')
    cvtp = {}
    for r in est.iter_rows(min_row=14, max_row=140, values_only=True):
        if r[0] in (None, '', 0):
            continue
        cvtp[(str(r[0]).upper(), str(r[1]).replace('.0', ''))] = r[9]       # J = CVTp %
    return dados, cvtp


def cv_lote(vals):
    if len(vals) < 2:
        return None
    m = st.mean(vals)
    return st.stdev(vals) / abs(m) * 100 if m else None


def pooled(lotes):
    num = den = 0.0
    for vals in lotes.values():
        if len(vals) >= N_LOTE_MIN:
            c = cv_lote(vals)
            if c is not None:
                num += (len(vals) - 1) * c * c
                den += len(vals) - 1
    return math.sqrt(num / den) if den else None


def main(produto, caminho, prefixo, meses=6):
    dados, cvtp = ler(caminho)
    fim = max(x['data'] for x in dados)
    ini = fim - timedelta(days=round(meses * 30.44)) + timedelta(days=1)
    grupos = defaultdict(lambda: defaultdict(list))
    for x in dados:
        if ini <= x['data'] <= fim:
            grupos[(x['eq'], x['an'], str(x['nv']).replace('.0', ''))][x['lote']].append(x)
    pares, suspeitos = [], []
    for (eq, an, nv), lotes in grupos.items():
        brutos = {l: [x['v'] for x in xs] for l, xs in lotes.items()}
        cv0 = pooled(brutos)
        limpos, sus = {}, []
        for l, xs in lotes.items():
            vs = [x['v'] for x in xs]
            med = st.median(vs)
            mad = st.median([abs(v - med) for v in vs]) * 1.4826 if len(vs) >= 5 else 0
            if mad == 0 and len(vs) >= 5:
                # mais da metade dos valores iguais (resultado inteiro, ex. Amilase 67 U/L): o MAD zera e um 482
                # passaria; reserva = desvio medio absoluto em torno da mediana (x 1,2533 ~ DP na normal)
                mad = st.mean([abs(v - med) for v in vs]) * 1.2533
            ok = []
            for x in xs:
                motivo = ''
                if x['v'] <= 0:
                    motivo = 'valor <= 0'
                elif mad > 0 and abs(x['v'] - med) / mad > Z_ROB:
                    motivo = f'{abs(x["v"] - med) / mad:.1f} DP robustos da mediana do lote ({med:.4g})'
                if motivo:
                    sus.append((x, l, motivo))
                else:
                    ok.append(x['v'])
            limpos[l] = ok
        cv1 = pooled(limpos)
        alvo = cvtp.get((an.upper(), nv))
        alvo = alvo if isinstance(alvo, (int, float)) else None
        n = sum(len(v) for v in brutos.values())
        lotes_ok = sum(1 for v in brutos.values() if len(v) >= N_LOTE_MIN)
        pares.append({'equipamento': eq, 'analito': an, 'nivel': nv, 'n': n, 'lotes_n20': lotes_ok,
                      'cv_pooled': cv0, 'cv_sem_suspeitos': cv1, 'cvtp': alvo, 'suspeitos': len(sus),
                      'acima_meta': (cv0 is not None and alvo is not None and cv0 > alvo)})
        for x, l, motivo in sus:
            suspeitos.append({'id': x['id'], 'equipamento': eq, 'analito': an, 'nivel': nv, 'lote': l,
                              'data_hora': x['dh'], 'resultado': x['v'], 'motivo': motivo})
    pares.sort(key=lambda p: -(p['cv_pooled'] or 0))
    with open(prefixo + '_pares.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=list(pares[0].keys()), delimiter=';')
        w.writeheader()
        w.writerows(pares)
    with open(prefixo + '_suspeitos.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=['id', 'equipamento', 'analito', 'nivel', 'lote', 'data_hora', 'resultado',
                                          'motivo'], delimiter=';')
        w.writeheader()
        w.writerows(sorted(suspeitos, key=lambda s: (s['analito'], s['nivel'], str(s['data_hora']))))
    acima = [p for p in pares if p['acima_meta']]
    print(f'{produto}: janela {ini:%d/%m/%Y}-{fim:%d/%m/%Y}; {len(pares)} pares; {len(acima)} com CV pooled > CVTp; '
          f'{len(suspeitos)} resultados suspeitos')
    return ini, fim, pares, suspeitos


if __name__ == '__main__':
    a = sys.argv[1:]
    m = int(a[a.index('--meses') + 1]) if '--meses' in a else 6
    main(a[0], a[1], a[2], m)
